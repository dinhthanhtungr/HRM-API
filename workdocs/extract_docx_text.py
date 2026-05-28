from pathlib import Path
from docx import Document


def iter_block_items(document):
    for paragraph in document.paragraphs:
        yield "P", paragraph.text

    for table_index, table in enumerate(document.tables, start=1):
        yield "T", f"[TABLE {table_index}]"
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            yield "R", " | ".join(cells)


source = Path("workdocs/Phong hop VietAus raw.docx")
output = Path("workdocs/Phong hop VietAus extracted.txt")
document = Document(source)

lines = []
for index, (kind, text) in enumerate(iter_block_items(document), start=1):
    if text.strip():
        lines.append(f"{index:04d} [{kind}] {text}")

output.write_text("\n".join(lines), encoding="utf-8")
print(output)
