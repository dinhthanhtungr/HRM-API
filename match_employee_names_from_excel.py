from __future__ import annotations

import argparse
import asyncio
import os
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any


APPSETTINGS_PATH = Path(r"F:\HRM.api\HRM.Api\appsettings.json")
DEFAULT_INPUT = Path(r"F:\Document\HRM - Mobile\DSNVV Gửi Tùng.xlsx")
DEFAULT_OUTPUT = Path(r"F:\Document\HRM - Mobile\DSNVV_Gui_Tung_employee_match.xlsx")


def _parse_npgsql_connection_string(value: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for part in value.split(";"):
        if not part.strip() or "=" not in part:
            continue
        key, raw_value = part.split("=", 1)
        normalized_key = key.strip().lower().replace(" ", "")
        result[normalized_key] = raw_value.strip()
    return result


def _read_default_db_config() -> dict[str, str] | None:
    database_url = os.getenv("DATABASE_URL", "").strip()
    if database_url:
        return {"dsn": database_url}

    if not APPSETTINGS_PATH.exists():
        return None

    text = APPSETTINGS_PATH.read_text(encoding="utf-8")
    match = re.search(r'"AppDbConnectionString"\s*:\s*"([^"]+)"', text)
    if not match:
        return None

    parts = _parse_npgsql_connection_string(match.group(1))
    return {
        "host": parts.get("host", "localhost"),
        "port": parts.get("port", "5432"),
        "database": parts.get("database", ""),
        "user": parts.get("username", parts.get("user", "")),
        "password": parts.get("password", ""),
    }


def _normalize_name(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d")
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _token_sort(value: str) -> str:
    return " ".join(sorted(value.split()))


def _name_score(left: Any, right: Any) -> float:
    left_norm = _normalize_name(left)
    right_norm = _normalize_name(right)
    if not left_norm or not right_norm:
        return 0.0

    if left_norm == right_norm:
        return 100.0

    direct_score = SequenceMatcher(None, left_norm, right_norm).ratio()
    token_score = SequenceMatcher(None, _token_sort(left_norm), _token_sort(right_norm)).ratio()
    left_tokens = set(left_norm.split())
    right_tokens = set(right_norm.split())
    overlap_score = 0.0
    if left_tokens and right_tokens:
        overlap_score = len(left_tokens & right_tokens) / max(len(left_tokens), len(right_tokens))

    return round(max(direct_score, token_score, overlap_score) * 100, 2)


async def _fetch_employees(db_config: dict[str, str] | None) -> list[dict[str, Any]]:
    try:
        import asyncpg
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Thieu thu vien asyncpg. Hay cai theo requirements.txt cua ETL_PO."
        ) from exc

    if not db_config:
        raise RuntimeError(
            "Khong tim thay cau hinh DB. Hay set DATABASE_URL hoac kiem tra HRM.Api/appsettings.json."
        )

    connect_kwargs: dict[str, Any]
    if "dsn" in db_config:
        connect_kwargs = {"dsn": db_config["dsn"]}
    else:
        connect_kwargs = {
            "host": db_config.get("host"),
            "port": int(db_config.get("port") or 5432),
            "database": db_config.get("database"),
            "user": db_config.get("user"),
            "password": db_config.get("password"),
        }

    query = """
        SELECT
            e."EmployeeID"::text AS "employeeId",
            e."ExternalId" AS "employeeExternalId",
            e."FullName" AS "employeeFullName",
            e."Status" AS "employeeStatus",
            e."IsActive" AS "employeeIsActive",
            u."Id"::text AS "aspNetUserId",
            u."UserName" AS "aspNetUserName",
            u."Email" AS "aspNetUserEmail",
            u."personName" AS "aspNetPersonName"
        FROM hr."Employees" e
        LEFT JOIN public."AspNetUsers" u
            ON u."EmployeeId" = e."EmployeeID";
    """

    conn = await asyncpg.connect(**connect_kwargs)
    try:
        records = await conn.fetch(query)
    finally:
        await conn.close()

    return [dict(record) for record in records]


def _best_matches(
    excel_name: Any,
    employees: list[dict[str, Any]],
    *,
    min_score: float,
    top: int,
) -> list[dict[str, Any]]:
    matches: list[dict[str, Any]] = []

    for employee in employees:
        employee_score = _name_score(excel_name, employee.get("employeeFullName"))
        person_score = _name_score(excel_name, employee.get("aspNetPersonName"))
        score = max(employee_score, person_score)
        if score < min_score:
            continue

        candidate = dict(employee)
        candidate["matchScore"] = score
        candidate["matchedBy"] = (
            "AspNetUsers.personName"
            if person_score > employee_score
            else "Employees.FullName"
        )
        candidate["hasAspNetUser"] = bool(employee.get("aspNetUserId"))
        matches.append(candidate)

    matches.sort(
        key=lambda item: (
            bool(item.get("hasAspNetUser")),
            float(item.get("matchScore") or 0),
            bool(item.get("employeeIsActive")),
        ),
        reverse=True,
    )
    return matches[:top]


def _write_result_workbook(
    *,
    input_path: Path,
    output_path: Path,
    sheet_name: str | None,
    sheet_index: int,
    name_col: int,
    header_row: int,
    employees: list[dict[str, Any]],
    min_score: float,
    top: int,
) -> None:
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Thieu thu vien openpyxl. Hay cai theo requirements.txt cua ETL_PO."
        ) from exc

    wb = load_workbook(input_path)
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.worksheets[sheet_index - 1]

    output_wb = load_workbook(input_path)
    if "Ket qua match" in output_wb.sheetnames:
        del output_wb["Ket qua match"]
    output_ws = output_wb.create_sheet("Ket qua match", 0)

    last_data_col = 1
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, "") and cell.column > last_data_col:
                last_data_col = cell.column

    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, last_data_col + 1):
            output_ws.cell(
                row=row_idx,
                column=col_idx,
                value=ws.cell(row=row_idx, column=col_idx).value,
            )

    appended_headers = [
        "MatchedEmployeeId",
        "MatchedEmployeeExternalId",
        "MatchedEmployeeFullName",
        "MatchedEmployeeStatus",
        "MatchedEmployeeIsActive",
        "AspNetUserId",
        "AspNetUserName",
        "AspNetUserEmail",
        "AspNetPersonName",
        "HasAspNetUser",
        "MatchScore",
        "MatchedBy",
        "OtherCandidates",
    ]

    start_col = last_data_col + 1
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for offset, header in enumerate(appended_headers):
        cell = output_ws.cell(row=header_row, column=start_col + offset, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx in range(header_row + 1, ws.max_row + 1):
        excel_name = ws.cell(row=row_idx, column=name_col).value
        matches = _best_matches(excel_name, employees, min_score=min_score, top=top)
        best = matches[0] if matches else {}
        other_candidates = []
        for candidate in matches[1:]:
            other_candidates.append(
                f'{candidate.get("employeeFullName")} | {candidate.get("employeeId")} | '
                f'score={candidate.get("matchScore")} | user={candidate.get("aspNetUserId") or ""}'
            )

        values = [
            best.get("employeeId"),
            best.get("employeeExternalId"),
            best.get("employeeFullName"),
            best.get("employeeStatus"),
            best.get("employeeIsActive"),
            best.get("aspNetUserId"),
            best.get("aspNetUserName"),
            best.get("aspNetUserEmail"),
            best.get("aspNetPersonName"),
            best.get("hasAspNetUser"),
            best.get("matchScore"),
            best.get("matchedBy"),
            "\n".join(other_candidates),
        ]

        for offset, value in enumerate(values):
            cell = output_ws.cell(row=row_idx, column=start_col + offset, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=offset in (2, 8, 12))

    output_ws.auto_filter.ref = output_ws.dimensions
    output_ws.freeze_panes = output_ws.cell(row=header_row + 1, column=1).coordinate

    for col_idx in range(start_col, start_col + len(appended_headers)):
        col_letter = get_column_letter(col_idx)
        output_ws.column_dimensions[col_letter].width = 22
    output_ws.column_dimensions[get_column_letter(start_col + 2)].width = 32
    output_ws.column_dimensions[get_column_letter(start_col + 8)].width = 28
    output_ws.column_dimensions[get_column_letter(start_col + 12)].width = 60

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)


async def run(args: argparse.Namespace) -> None:
    db_config = _read_default_db_config()
    if args.db_host or args.db_name or args.db_user or args.db_password or args.db_port:
        db_config = {
            "host": args.db_host,
            "port": str(args.db_port or 5432),
            "database": args.db_name,
            "user": args.db_user,
            "password": args.db_password,
        }

    employees = await _fetch_employees(db_config)
    _write_result_workbook(
        input_path=args.input,
        output_path=args.output,
        sheet_name=args.sheet_name,
        sheet_index=args.sheet_index,
        name_col=args.name_col,
        header_row=args.header_row,
        employees=employees,
        min_score=args.min_score,
        top=args.top,
    )
    print(f"Da doc {len(employees)} nhan vien tu DB.")
    print(f"Da xuat file: {args.output}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description='Match cot Ho ten trong Excel voi hr."Employees" va uu tien nhan vien co public."AspNetUsers".'
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--sheet-name", default="Danh sách nhân viên")
    parser.add_argument("--sheet-index", type=int, default=2)
    parser.add_argument("--name-col", type=int, default=4, help="Cot Ho ten trong Excel, mac dinh la cot 4.")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--min-score", type=float, default=75.0)
    parser.add_argument("--top", type=int, default=5)
    parser.add_argument("--db-host", default="")
    parser.add_argument("--db-port", type=int, default=0)
    parser.add_argument("--db-name", default="")
    parser.add_argument("--db-user", default="")
    parser.add_argument("--db-password", default="")
    return parser.parse_args()


def main() -> None:
    asyncio.run(run(parse_args()))


if __name__ == "__main__":
    main()
