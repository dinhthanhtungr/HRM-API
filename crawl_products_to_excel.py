from __future__ import annotations

import argparse
import asyncio
import math
import os
import re
import sys
import uuid
from pathlib import Path
from typing import Any


ETL_DIR = Path(r"F:\crawlData\CawlDataForVAE\ETL_PO")
APPSETTINGS_PATH = Path(r"F:\HRM.api\HRM.Api\appsettings.json")
DEFAULT_BASE_URL = "http://192.168.7.4"
DEFAULT_OUTPUT = ETL_DIR / "Data" / "products_comment_rosh_check.xlsx"
SHEET_NAME = "Products"
API_FIELDS = ("id", "externalId", "comment", "roshStandard", "colourCode", "status")
DB_FIELDS = (
    "dbProductId",
    "dbRohsStandard",
    "dbSampleRequestId",
    "dbSampleRequestExternalId",
    "dbSaleComment",
    "dbSampleRequestStatus",
)
PREVIEW_FIELDS = (
    "productId",
    "productExternalId",
    "colourCode",
    "apiStatus",
    "dbProductFound",
    "currentProductRohsStandard",
    "apiRoshStandard",
    "willChangeProductRohsStandard",
    "sampleRequestFound",
    "sampleRequestId",
    "sampleRequestExternalId",
    "sampleRequestStatus",
    "currentSaleComment",
    "apiComment",
    "willChangeSaleComment",
)


def _prepare_auth_import() -> None:
    """Load ETL .env and make the existing auth_token.py importable."""
    try:
        from dotenv import load_dotenv
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Thieu thu vien python-dotenv. Hay cai theo requirements.txt cua ETL_PO."
        ) from exc

    if not ETL_DIR.exists():
        raise RuntimeError(f"Khong tim thay thu muc ETL: {ETL_DIR}")

    env_path = ETL_DIR / ".env"
    if env_path.exists():
        load_dotenv(env_path, override=False)

    etl_dir_text = str(ETL_DIR)
    if etl_dir_text not in sys.path:
        sys.path.insert(0, etl_dir_text)


def _safe_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _is_different(current_value: Any, next_value: Any) -> bool:
    return current_value != next_value


def _parse_npgsql_connection_string(value: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for part in value.split(";"):
        if not part.strip() or "=" not in part:
            continue
        key, raw_value = part.split("=", 1)
        normalized_key = key.strip().lower().replace(" ", "")
        result[normalized_key] = raw_value.strip()
    return result


def _read_default_db_config() -> dict[str, str] | None:
    database_url = os.getenv("DATABASE_URL", "").strip()
    if database_url:
        return {"dsn": database_url}

    if not APPSETTINGS_PATH.exists():
        return None

    text = APPSETTINGS_PATH.read_text(encoding="utf-8")
    match = re.search(r'"AppDbConnectionString"\s*:\s*"([^"]+)"', text)
    if not match:
        return None

    parts = _parse_npgsql_connection_string(match.group(1))
    return {
        "host": parts.get("host", "localhost"),
        "port": parts.get("port", "5432"),
        "database": parts.get("database", ""),
        "user": parts.get("username", parts.get("user", "")),
        "password": parts.get("password", ""),
    }


def _extract_products(payload: dict[str, Any]) -> list[dict[str, Any]]:
    products = payload.get("product")
    if isinstance(products, list):
        return [item for item in products if isinstance(item, dict)]
    return []


def _extract_total(payload: dict[str, Any]) -> int | None:
    meta = payload.get("meta")
    if not isinstance(meta, dict):
        return None

    total = meta.get("total")
    if isinstance(total, int):
        return total
    if isinstance(total, str) and total.isdigit():
        return int(total)
    return None


async def _fetch_products_page(
    client: Any,
    *,
    base_url: str,
    headers: dict[str, str],
    page_number: int,
    page_size: int,
    external_id: str,
    status: str,
) -> dict[str, Any]:
    url = f"{base_url.rstrip('/')}/api/api/products"
    params = {
        "externalId": external_id,
        "pageNumber": page_number,
        "pageSize": page_size,
        "status": status,
    }

    response = await client.get(url, params=params, headers=headers)
    response.raise_for_status()
    return response.json()


async def _fetch_db_product_sample_request_map(
    product_ids: list[str],
    *,
    db_config: dict[str, str] | None,
) -> dict[str, list[dict[str, Any]]]:
    if not product_ids:
        return {}

    try:
        import asyncpg
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Thieu thu vien asyncpg. Hay cai theo requirements.txt cua ETL_PO."
        ) from exc

    if not db_config:
        raise RuntimeError(
            "Khong tim thay cau hinh DB. Hay set DATABASE_URL hoac kiem tra HRM.Api/appsettings.json."
        )

    unique_ids: list[uuid.UUID] = []
    seen_ids: set[uuid.UUID] = set()
    for product_id in product_ids:
        try:
            parsed_id = uuid.UUID(str(product_id))
        except ValueError:
            continue
        if parsed_id not in seen_ids:
            seen_ids.add(parsed_id)
            unique_ids.append(parsed_id)

    if not unique_ids:
        return {}
    query = """
        SELECT
            p."ProductId"::text AS "dbProductId",
            p."RohsStandard" AS "dbRohsStandard",
            sr."SampleRequestId"::text AS "dbSampleRequestId",
            sr."ExternalId" AS "dbSampleRequestExternalId",
            sr."SaleComment" AS "dbSaleComment",
            sr."Status" AS "dbSampleRequestStatus"
        FROM "SampleRequests"."Products" p
        LEFT JOIN "SampleRequests"."SampleRequests" sr
            ON sr."ProductId" = p."ProductId"
           AND sr."IsActive" = TRUE
        WHERE p."ProductId" = ANY($1::uuid[])
        ORDER BY p."ProductId", sr."CreatedDate" DESC NULLS LAST, sr."SampleRequestId" DESC NULLS LAST;
    """

    connect_kwargs: dict[str, Any]
    if "dsn" in db_config:
        connect_kwargs = {"dsn": db_config["dsn"]}
    else:
        connect_kwargs = {
            "host": db_config.get("host"),
            "port": int(db_config.get("port") or 5432),
            "database": db_config.get("database"),
            "user": db_config.get("user"),
            "password": db_config.get("password"),
        }

    conn = await asyncpg.connect(**connect_kwargs)
    try:
        records = await conn.fetch(query, unique_ids)
    finally:
        await conn.close()

    result: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        result.setdefault(str(record["dbProductId"]), []).append(dict(record))
    return result


async def _apply_api_updates_to_db(
    api_rows: list[dict[str, Any]],
    *,
    db_config: dict[str, str] | None,
) -> tuple[int, int]:
    try:
        import asyncpg
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Thieu thu vien asyncpg. Hay cai theo requirements.txt cua ETL_PO."
        ) from exc

    if not db_config:
        raise RuntimeError(
            "Khong tim thay cau hinh DB. Hay set DATABASE_URL hoac kiem tra HRM.Api/appsettings.json."
        )

    update_items: list[tuple[uuid.UUID, Any, Any]] = []
    seen_ids: set[uuid.UUID] = set()
    for row in api_rows:
        try:
            product_id = uuid.UUID(str(row.get("id") or ""))
        except ValueError:
            continue

        if product_id in seen_ids:
            continue
        seen_ids.add(product_id)

        update_items.append(
            (
                product_id,
                row.get("roshStandard"),
                row.get("comment") if "comment" in row else None,
            )
        )

    if not update_items:
        return 0, 0

    connect_kwargs: dict[str, Any]
    if "dsn" in db_config:
        connect_kwargs = {"dsn": db_config["dsn"]}
    else:
        connect_kwargs = {
            "host": db_config.get("host"),
            "port": int(db_config.get("port") or 5432),
            "database": db_config.get("database"),
            "user": db_config.get("user"),
            "password": db_config.get("password"),
        }

    product_sql = """
        UPDATE "SampleRequests"."Products" AS p
        SET "RohsStandard" = v."RohsStandard"
        FROM (
            SELECT
                unnest($1::uuid[]) AS "ProductId",
                unnest($2::boolean[]) AS "RohsStandard"
        ) AS v
        WHERE p."ProductId" = v."ProductId"
          AND p."RohsStandard" IS DISTINCT FROM v."RohsStandard";
    """
    sample_request_sql = """
        UPDATE "SampleRequests"."SampleRequests" AS sr
        SET "SaleComment" = v."SaleComment"
        FROM (
            SELECT
                unnest($1::uuid[]) AS "ProductId",
                unnest($2::text[]) AS "SaleComment"
        ) AS v
        WHERE sr."ProductId" = v."ProductId"
          AND sr."IsActive" = TRUE
          AND sr."SaleComment" IS DISTINCT FROM v."SaleComment";
    """

    rohs_items = [
        (product_id, rohs)
        for product_id, rohs, _ in update_items
        if isinstance(rohs, bool)
    ]
    comment_items = [
        (product_id, comment)
        for product_id, _, comment in update_items
    ]

    conn = await asyncpg.connect(**connect_kwargs)
    try:
        async with conn.transaction():
            product_result = "UPDATE 0"
            sample_request_result = "UPDATE 0"

            if rohs_items:
                product_result = await conn.execute(
                    product_sql,
                    [item[0] for item in rohs_items],
                    [item[1] for item in rohs_items],
                )

            if comment_items:
                sample_request_result = await conn.execute(
                    sample_request_sql,
                    [item[0] for item in comment_items],
                    [item[1] for item in comment_items],
                )
    finally:
        await conn.close()

    def parse_update_count(result: str) -> int:
        parts = result.split()
        if len(parts) == 2 and parts[0] == "UPDATE" and parts[1].isdigit():
            return int(parts[1])
        return 0

    return parse_update_count(product_result), parse_update_count(sample_request_result)


def _build_preview_rows(
    api_rows: list[dict[str, Any]],
    db_map: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    preview_rows: list[dict[str, Any]] = []

    for api_row in api_rows:
        product_id = str(api_row.get("id") or "")
        db_rows = db_map.get(product_id)

        if not db_rows:
            db_rows = [{}]

        for db_row in db_rows:
            db_product_found = bool(db_row.get("dbProductId"))
            sample_request_found = bool(db_row.get("dbSampleRequestId"))
            current_rohs = db_row.get("dbRohsStandard")
            next_rohs = api_row.get("roshStandard")
            current_sale_comment = db_row.get("dbSaleComment")
            next_sale_comment = api_row.get("comment")

            preview_rows.append(
                {
                    "productId": product_id,
                    "productExternalId": api_row.get("externalId"),
                    "colourCode": api_row.get("colourCode"),
                    "apiStatus": api_row.get("status"),
                    "dbProductFound": db_product_found,
                    "currentProductRohsStandard": current_rohs,
                    "apiRoshStandard": next_rohs,
                    "willChangeProductRohsStandard": db_product_found
                    and _is_different(current_rohs, next_rohs),
                    "sampleRequestFound": sample_request_found,
                    "sampleRequestId": db_row.get("dbSampleRequestId"),
                    "sampleRequestExternalId": db_row.get("dbSampleRequestExternalId"),
                    "sampleRequestStatus": db_row.get("dbSampleRequestStatus"),
                    "currentSaleComment": current_sale_comment,
                    "apiComment": next_sale_comment,
                    "willChangeSaleComment": sample_request_found
                    and _is_different(current_sale_comment, next_sale_comment),
                }
            )

    return preview_rows


def _write_excel(rows: list[dict[str, Any]], output_path: Path, fields: tuple[str, ...]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Thieu thu vien openpyxl. Hay cai theo requirements.txt cua ETL_PO."
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(list(fields))

    for row in rows:
        ws.append([_safe_cell_value(row.get(field)) for field in fields])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    width_by_col = {
        "A": 40,
        "B": 18,
        "C": 16,
        "D": 14,
        "E": 16,
        "F": 16,
        "G": 16,
        "H": 16,
        "I": 16,
        "J": 40,
        "K": 22,
        "L": 20,
        "M": 70,
        "N": 70,
        "O": 18,
    }
    for col, width in width_by_col.items():
        ws.column_dimensions[col].width = width

    wrap_columns = {"C", "M", "N"}
    if "comment" in fields:
        wrap_columns.add(get_column_letter(fields.index("comment") + 1))
    if "dbSaleComment" in fields:
        wrap_columns.add(get_column_letter(fields.index("dbSaleComment") + 1))

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                wrap_text=cell.column_letter in wrap_columns,
                vertical="top",
            )

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 36

    for col_idx in range(1, len(fields) + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].bestFit = True

    wb.save(output_path)


async def crawl_products_to_excel(
    *,
    base_url: str,
    output_path: Path,
    page_size: int,
    max_pages: int | None,
    external_id: str,
    status: str,
    include_db: bool,
    db_config: dict[str, str] | None,
    apply_update: bool,
) -> None:
    try:
        import httpx
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Thieu thu vien httpx. Hay cai theo requirements.txt cua ETL_PO."
        ) from exc

    _prepare_auth_import()
    from auth_token import force_refresh_auth_headers, get_auth_headers

    headers = await get_auth_headers()
    rows: list[dict[str, Any]] = []

    async with httpx.AsyncClient(timeout=httpx.Timeout(60.0, connect=15.0)) as client:
        async def fetch_page_with_refresh(page_number: int) -> dict[str, Any]:
            nonlocal headers
            try:
                return await _fetch_products_page(
                    client,
                    base_url=base_url,
                    headers=headers,
                    page_number=page_number,
                    page_size=page_size,
                    external_id=external_id,
                    status=status,
                )
            except httpx.HTTPStatusError as exc:
                if exc.response.status_code not in (401, 403):
                    raise
                headers = await force_refresh_auth_headers()
                return await _fetch_products_page(
                    client,
                    base_url=base_url,
                    headers=headers,
                    page_number=page_number,
                    page_size=page_size,
                    external_id=external_id,
                    status=status,
                )

        first_payload = await fetch_page_with_refresh(1)

        rows.extend(_extract_products(first_payload))
        total = _extract_total(first_payload)
        total_pages = math.ceil(total / page_size) if total else None
        pages_to_fetch = total_pages or 1
        if max_pages is not None:
            pages_to_fetch = min(pages_to_fetch, max_pages)

        print(
            f"Page 1/{pages_to_fetch}: lay {len(rows)} dong"
            + (f" tren tong {total} dong" if total else "")
        )

        for page_number in range(2, pages_to_fetch + 1):
            payload = await fetch_page_with_refresh(page_number)
            page_rows = _extract_products(payload)
            rows.extend(page_rows)
            print(f"Page {page_number}/{pages_to_fetch}: lay them {len(page_rows)} dong")

    fields = API_FIELDS
    export_rows = [{field: row.get(field) for field in API_FIELDS} for row in rows]

    if include_db:
        product_ids = [str(row.get("id")) for row in rows if row.get("id")]
        db_map = await _fetch_db_product_sample_request_map(product_ids, db_config=db_config)
        export_rows = _build_preview_rows(rows, db_map)
        fields = PREVIEW_FIELDS

    _write_excel(export_rows, output_path, fields)
    print(f"Da ghi {len(export_rows)} dong vao: {output_path}")

    if apply_update:
        product_count, sample_request_count = await _apply_api_updates_to_db(
            rows,
            db_config=db_config,
        )
        print(
            "Da update DB: "
            f"{product_count} dong Products.RohsStandard, "
            f"{sample_request_count} dong SampleRequests.SaleComment"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Cao product tu API va luu cac cot id/externalId/comment/roshStandard/colourCode/status vao Excel."
    )
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--page-size", type=int, default=100)
    parser.add_argument("--max-pages", type=int, default=3, help="Mac dinh chi cao 3 trang de kiem tra. Dung 0 de cao tat ca.")
    parser.add_argument("--external-id", default="")
    parser.add_argument("--status", default="")
    parser.add_argument(
        "--include-db",
        action="store_true",
        help="Doc DB va xuat preview: hien tai dang la gi, neu update tu API thi thanh gi.",
    )
    parser.add_argument("--db-host", default="")
    parser.add_argument("--db-port", type=int, default=0)
    parser.add_argument("--db-name", default="")
    parser.add_argument("--db-user", default="")
    parser.add_argument("--db-password", default="")
    parser.add_argument(
        "--apply-update",
        action="store_true",
        help="Thuc su update DB. Chi update Products.RohsStandard va SampleRequests.SaleComment.",
    )
    parser.add_argument(
        "--confirm-update",
        default="",
        help='Can nhap dung chu "UPDATE_2_FIELDS" khi dung --apply-update.',
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.apply_update and args.confirm_update != "UPDATE_2_FIELDS":
        raise SystemExit(
            'De update DB, chay them: --confirm-update UPDATE_2_FIELDS'
        )

    max_pages = None if args.max_pages == 0 else args.max_pages
    db_config = _read_default_db_config()
    if args.db_host or args.db_name or args.db_user or args.db_password or args.db_port:
        db_config = {
            "host": args.db_host,
            "port": str(args.db_port or 5432),
            "database": args.db_name,
            "user": args.db_user,
            "password": args.db_password,
        }

    asyncio.run(
        crawl_products_to_excel(
            base_url=args.base_url,
            output_path=args.output,
            page_size=args.page_size,
            max_pages=max_pages,
            external_id=args.external_id,
            status=args.status,
            include_db=args.include_db or args.apply_update,
            db_config=db_config,
            apply_update=args.apply_update,
        )
    )


if __name__ == "__main__":
    main()
