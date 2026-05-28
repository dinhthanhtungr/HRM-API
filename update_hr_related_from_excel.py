from __future__ import annotations

import argparse
import asyncio
import os
import re
import unicodedata
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any


APPSETTINGS_PATH = Path(r"F:\HRM.api\HRM.Api\appsettings.json")
DEFAULT_INPUT = Path(r"F:\Document\HRM - Mobile\DSNVV_Gui_Tung_employee_match_isAllowChange.xlsx")
DEFAULT_LOG = Path(r"F:\Document\HRM - Mobile\HR_related_update_preview.xlsx")

SHEET_MATCH_INDEX = 0
SHEET_DATA_INDEX = 2
MATCH_HEADER_ROW = 1
DATA_HEADER_ROW = 2
DATA_START_ROW = 3


def _parse_npgsql_connection_string(value: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for part in value.split(";"):
        if not part.strip() or "=" not in part:
            continue
        key, raw_value = part.split("=", 1)
        result[key.strip().lower().replace(" ", "")] = raw_value.strip()
    return result


def _read_default_db_config() -> dict[str, str] | None:
    database_url = os.getenv("DATABASE_URL", "").strip()
    if database_url:
        return {"dsn": database_url}

    if not APPSETTINGS_PATH.exists():
        return None

    text = APPSETTINGS_PATH.read_text(encoding="utf-8")
    match = re.search(r'"AppDbConnectionString"\s*:\s*"([^"]+)"', text)
    if not match:
        return None

    parts = _parse_npgsql_connection_string(match.group(1))
    return {
        "host": parts.get("host", "localhost"),
        "port": parts.get("port", "5432"),
        "database": parts.get("database", ""),
        "user": parts.get("username", parts.get("user", "")),
        "password": parts.get("password", ""),
    }


def _normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFD", text.strip().lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d")
    return re.sub(r"[^a-z0-9]+", "", text)


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _text(value: Any) -> str | None:
    if _is_blank(value):
        return None
    text = str(value).strip()
    return text if text.lower() != "none" else None


def _date(value: Any) -> date | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _date_from_ymd(year_value: Any, month_value: Any, day_value: Any) -> date | None:
    try:
        year = int(float(str(year_value).strip()))
        month = int(float(str(month_value).strip()))
        day = int(float(str(day_value).strip()))
        return date(year, month, day)
    except Exception:
        return None


def _bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    return text in {"true", "1", "yes", "y", "x", "co", "có"}


def _education_level(value: Any) -> int | None:
    text = _normalize_header(value)
    if not text:
        return None
    if "daihoc" in text or "cunhan" in text:
        return 6
    if "caodang" in text:
        return 5
    if "trungcap" in text or "nghe" in text:
        return 4
    if "thpt" in text or "12" in text:
        return 3
    if "thcs" in text:
        return 2
    if "tieuhoc" in text:
        return 1
    if "thacsi" in text:
        return 7
    if "tiensi" in text:
        return 8
    return 99


def _contract_type(value: Any) -> int | None:
    text = _normalize_header(value)
    if not text:
        return None
    if "thu viec" in text or "thuviec" in text:
        return 1
    if "vth" in text or "khongthoihan" in text:
        return 3
    if "12" in text or "24" in text or "hdl" in text or "hdld" in text:
        return 2
    return 99


def _connect_kwargs(db_config: dict[str, str] | None) -> dict[str, Any]:
    if not db_config:
        raise RuntimeError("Khong tim thay cau hinh DB.")
    if "dsn" in db_config:
        return {"dsn": db_config["dsn"]}
    return {
        "host": db_config.get("host"),
        "port": int(db_config.get("port") or 5432),
        "database": db_config.get("database"),
        "user": db_config.get("user"),
        "password": db_config.get("password"),
    }


def _header_map(ws: Any, header_row: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        key = _normalize_header(value)
        if key and key not in result:
            result[key] = col
    return result


def _find_col(headers: dict[str, int], *names: str) -> int | None:
    for name in names:
        col = headers.get(_normalize_header(name))
        if col:
            return col
    return None


def _cell(ws: Any, row: int, col: int | None) -> Any:
    return ws.cell(row, col).value if col else None


def _load_allowed_rows(input_path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(input_path, data_only=True)
    match_ws = wb.worksheets[SHEET_MATCH_INDEX]
    data_ws = wb.worksheets[SHEET_DATA_INDEX]

    match_headers = _header_map(match_ws, MATCH_HEADER_ROW)
    data_headers = _header_map(data_ws, DATA_HEADER_ROW)

    employee_id_col = _find_col(match_headers, "MatchedEmployeeId")
    allow_col = _find_col(match_headers, "isAllowChange")
    if not employee_id_col or not allow_col:
        raise RuntimeError("Khong tim thay MatchedEmployeeId hoac isAllowChange trong sheet 1.")

    cols = {
        "attendance_code": _find_col(data_headers, "Mã chấm công"),
        "full_name": _find_col(data_headers, "Họ tên"),
        "ethnicity": _find_col(data_headers, "Dân tộc"),
        "date_hired": _find_col(data_headers, "Thời gian vào"),
        "probation_end_date": _find_col(data_headers, "Ngày kết thúc thử việc"),
        "onboarding_training_date": _find_col(data_headers, "Ngày đào tạo hội nhập"),
        "contract_type": _find_col(data_headers, "Loại hợp đồng"),
        "contract_12": _find_col(data_headers, "HĐLĐ 12 tháng"),
        "contract_24": _find_col(data_headers, "HĐLĐ 24 tháng"),
        "contract_vth": _find_col(data_headers, "HDLĐ VTH"),
        "identifier_issue_date": _find_col(data_headers, "Ngày cấp"),
        "identifier_issue_place": _find_col(data_headers, "Nơi cấp"),
        "education_level": _find_col(data_headers, "Trình độ học vấn"),
        "permanent_address": _find_col(data_headers, "Địa chỉ thường trú"),
        "temporary_address": _find_col(data_headers, "Địa chỉ tạm trú"),
        "tax_code": _find_col(data_headers, "MST"),
        "bank_scb": _find_col(data_headers, "TK SCB"),
        "bank_other": _find_col(data_headers, "TK KHAC"),
        "bank_name": _find_col(data_headers, "TÊN NH"),
        "work_group_text": _find_col(data_headers, "Nhóm"),
        "relative_full_name": _find_col(data_headers, "Người thân"),
    }

    rows: list[dict[str, Any]] = []
    for row_idx in range(DATA_START_ROW, data_ws.max_row + 1):
        if not _bool(match_ws.cell(row_idx, allow_col).value):
            continue

        employee_id_text = _text(match_ws.cell(row_idx, employee_id_col).value)
        if not employee_id_text:
            continue

        try:
            employee_id = uuid.UUID(employee_id_text)
        except ValueError:
            continue

        contract_type_raw = _cell(data_ws, row_idx, cols["contract_type"])
        contract_start = (
            _date(_cell(data_ws, row_idx, cols["contract_vth"]))
            or _date(_cell(data_ws, row_idx, cols["contract_24"]))
            or _date(_cell(data_ws, row_idx, cols["contract_12"]))
            or _date(_cell(data_ws, row_idx, cols["date_hired"]))
        )

        bank_account = _text(_cell(data_ws, row_idx, cols["bank_scb"])) or _text(
            _cell(data_ws, row_idx, cols["bank_other"])
        )

        rows.append(
            {
                "excel_row": row_idx,
                "employee_id": employee_id,
                "full_name": _text(_cell(data_ws, row_idx, cols["full_name"])),
                "attendance_code": _text(_cell(data_ws, row_idx, cols["attendance_code"])),
                "ethnicity": _text(_cell(data_ws, row_idx, cols["ethnicity"])),
                "education_level": _education_level(_cell(data_ws, row_idx, cols["education_level"])),
                "identifier_issue_date": _date(_cell(data_ws, row_idx, cols["identifier_issue_date"])),
                "identifier_issue_place": _text(_cell(data_ws, row_idx, cols["identifier_issue_place"])),
                "permanent_address": _text(_cell(data_ws, row_idx, cols["permanent_address"])),
                "temporary_address": _text(_cell(data_ws, row_idx, cols["temporary_address"])),
                "tax_code": _text(_cell(data_ws, row_idx, cols["tax_code"])),
                "bank_account": bank_account,
                "bank_name": _text(_cell(data_ws, row_idx, cols["bank_name"])),
                "account_holder": _text(_cell(data_ws, row_idx, cols["full_name"])),
                "work_location": _text(_cell(data_ws, row_idx, cols["work_group_text"])),
                "probation_end_date": _date(_cell(data_ws, row_idx, cols["probation_end_date"])),
                "onboarding_training_date": _date(_cell(data_ws, row_idx, cols["onboarding_training_date"])),
                "relative_full_name": _text(_cell(data_ws, row_idx, cols["relative_full_name"])),
                "contract_type": _contract_type(contract_type_raw),
                "contract_start_date": contract_start,
            }
        )
    return rows


async def _existing_employee_ids(conn: Any, rows: list[dict[str, Any]]) -> set[uuid.UUID]:
    ids = list({row["employee_id"] for row in rows})
    records = await conn.fetch(
        'SELECT "EmployeeID"::text FROM hr."Employees" WHERE "EmployeeID" = ANY($1::uuid[])',
        ids,
    )
    return {uuid.UUID(record["EmployeeID"]) for record in records}


async def _apply_updates(conn: Any, rows: list[dict[str, Any]]) -> dict[str, int]:
    counts = {
        "employee_profiles": 0,
        "employee_insurance_profiles": 0,
        "employee_bank_accounts": 0,
        "employee_work_profiles": 0,
        "employee_relatives": 0,
        "employee_contracts": 0,
    }

    profile_sql = """
        INSERT INTO hr.employee_profiles (
            employee_profile_id, employee_id, ethnicity, education_level,
            identifier_issue_date, identifier_issue_place, permanent_address, temporary_address
        )
        VALUES (gen_random_uuid(), $1, $2, $3, $4, $5, $6, $7)
        ON CONFLICT (employee_id) DO UPDATE SET
            ethnicity = EXCLUDED.ethnicity,
            education_level = EXCLUDED.education_level,
            identifier_issue_date = EXCLUDED.identifier_issue_date,
            identifier_issue_place = EXCLUDED.identifier_issue_place,
            permanent_address = EXCLUDED.permanent_address,
            temporary_address = EXCLUDED.temporary_address;
    """
    insurance_sql = """
        INSERT INTO hr.employee_insurance_profiles (
            employee_insurance_profile_id, employee_id, tax_code
        )
        VALUES (gen_random_uuid(), $1, $2)
        ON CONFLICT (employee_id) DO UPDATE SET
            tax_code = EXCLUDED.tax_code;
    """
    work_profile_sql = """
        UPDATE hr.employee_work_profiles
        SET
            attendance_code = COALESCE($2, attendance_code),
            work_location = COALESCE($3, work_location),
            probation_end_date = COALESCE($4, probation_end_date),
            onboarding_training_date = COALESCE($5, onboarding_training_date)
        WHERE employee_id = $1
          AND is_current = TRUE
          AND is_active = TRUE;
    """
    bank_sql = """
        INSERT INTO hr.employee_bank_accounts (
            employee_bank_account_id, employee_id, bank_name, account_number, account_holder, is_payroll_account
        )
        VALUES (gen_random_uuid(), $1, $2, $3, $4, TRUE)
        ON CONFLICT DO NOTHING;
    """
    update_bank_sql = """
        UPDATE hr.employee_bank_accounts
        SET bank_name = COALESCE($2, bank_name),
            account_number = $3,
            account_holder = COALESCE($4, account_holder),
            is_payroll_account = TRUE
        WHERE employee_bank_account_id = (
            SELECT employee_bank_account_id
            FROM hr.employee_bank_accounts
            WHERE employee_id = $1
            ORDER BY is_payroll_account DESC, employee_bank_account_id
            LIMIT 1
        );
    """
    relative_sql = """
        INSERT INTO hr.employee_relatives (
            employee_relative_id, employee_id, full_name, relationship, phone_number, is_emergency_contact
        )
        VALUES (gen_random_uuid(), $1, $2, 99, NULL, TRUE);
    """
    contract_sql = """
        UPDATE hr.employee_contracts
        SET
            contract_type = COALESCE($2, contract_type),
            start_date = COALESCE($3, start_date),
            is_current = TRUE
        WHERE employee_id = $1
          AND is_current = TRUE;
    """

    for row in rows:
        employee_id = row["employee_id"]

        await conn.execute(
            profile_sql,
            employee_id,
            row["ethnicity"],
            row["education_level"],
            row["identifier_issue_date"],
            row["identifier_issue_place"],
            row["permanent_address"],
            row["temporary_address"],
        )
        counts["employee_profiles"] += 1

        if row["tax_code"]:
            await conn.execute(insurance_sql, employee_id, row["tax_code"])
            counts["employee_insurance_profiles"] += 1

        result = await conn.execute(
            work_profile_sql,
            employee_id,
            row["attendance_code"],
            row["work_location"],
            row["probation_end_date"],
            row["onboarding_training_date"],
        )
        counts["employee_work_profiles"] += int(result.split()[-1])

        if row["bank_account"]:
            result = await conn.execute(
                update_bank_sql,
                employee_id,
                row["bank_name"],
                row["bank_account"],
                row["account_holder"],
            )
            updated = int(result.split()[-1])
            if not updated:
                await conn.execute(
                    bank_sql,
                    employee_id,
                    row["bank_name"],
                    row["bank_account"],
                    row["account_holder"],
                )
                updated = 1
            counts["employee_bank_accounts"] += updated

        if row["relative_full_name"]:
            exists = await conn.fetchval(
                """
                SELECT 1
                FROM hr.employee_relatives
                WHERE employee_id = $1 AND lower(full_name) = lower($2)
                LIMIT 1
                """,
                employee_id,
                row["relative_full_name"],
            )
            if not exists:
                await conn.execute(relative_sql, employee_id, row["relative_full_name"])
                counts["employee_relatives"] += 1

        if row["contract_type"] and row["contract_start_date"]:
            result = await conn.execute(
                contract_sql,
                employee_id,
                row["contract_type"],
                row["contract_start_date"],
            )
            counts["employee_contracts"] += int(result.split()[-1])

    return counts


def _write_log(rows: list[dict[str, Any]], output_path: Path, *, apply_update: bool) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "HR related update"
    headers = [
        "mode",
        "excel_row",
        "employee_id",
        "full_name",
        "attendance_code",
        "ethnicity",
        "education_level",
        "identifier_issue_date",
        "identifier_issue_place",
        "permanent_address",
        "temporary_address",
        "tax_code",
        "bank_account",
        "bank_name",
        "work_location",
        "probation_end_date",
        "onboarding_training_date",
        "relative_full_name",
        "contract_type",
        "contract_start_date",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)

    mode = "UPDATED" if apply_update else "PREVIEW_ONLY"
    for row in rows:
        ws.append([mode] + [str(row.get(header, "") or "") for header in headers[1:]])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 22
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


async def run(args: argparse.Namespace) -> None:
    if args.apply_update and args.confirm_update != "UPDATE_HR_RELATED":
        raise SystemExit('Muon update that, them: --confirm-update UPDATE_HR_RELATED')

    rows = _load_allowed_rows(args.input)
    _write_log(rows, args.log, apply_update=args.apply_update)

    db_config = _read_default_db_config()
    if args.db_host or args.db_name or args.db_user or args.db_password or args.db_port:
        db_config = {
            "host": args.db_host,
            "port": str(args.db_port or 5432),
            "database": args.db_name,
            "user": args.db_user,
            "password": args.db_password,
        }

    if not args.apply_update:
        print(f"PREVIEW: {len(rows)} dong allow TRUE se duoc xu ly.")
        print(f"Da xuat log preview: {args.log}")
        return

    try:
        import asyncpg
    except ModuleNotFoundError as exc:
        raise RuntimeError("Thieu asyncpg. Hay cai dependencies truoc.") from exc

    conn = await asyncpg.connect(**_connect_kwargs(db_config))
    try:
        existing_ids = await _existing_employee_ids(conn, rows)
        rows = [row for row in rows if row["employee_id"] in existing_ids]
        async with conn.transaction():
            counts = await _apply_updates(conn, rows)
    finally:
        await conn.close()

    _write_log(rows, args.log, apply_update=True)
    print(f"UPDATED: {len(rows)} nhan vien allow TRUE co EmployeeId ton tai.")
    for table, count in counts.items():
        print(f"{table}: {count}")
    print(f"Da xuat log update: {args.log}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Doc DSNVV_Gui_Tung_employee_match_isAllowChange.xlsx va cap nhat cac bang HR lien quan "
            "employee, khong update hr.Employees."
        )
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--log", type=Path, default=DEFAULT_LOG)
    parser.add_argument("--apply-update", action="store_true")
    parser.add_argument("--confirm-update", default="")
    parser.add_argument("--db-host", default="")
    parser.add_argument("--db-port", type=int, default=0)
    parser.add_argument("--db-name", default="")
    parser.add_argument("--db-user", default="")
    parser.add_argument("--db-password", default="")
    return parser.parse_args()


def main() -> None:
    asyncio.run(run(parse_args()))


if __name__ == "__main__":
    main()
